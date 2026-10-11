"""Run with the existing API container's manage.py shell; creates only disposable fixtures."""

import gzip
import hashlib
import json
import os
import random
import shutil
import sys
import time
import traceback
import uuid
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Barrier

import httpx
import pyarrow as pa
import pyarrow.parquet as pq
import pypdfium2 as pdfium
import requests
from django.db import transaction
from django.utils import timezone
from docx import Document
from jsonschema import Draft202012Validator

from overbae.models import (
    APIToken,
    BillingTelemetry,
    Capability,
    Cell,
    DataExploration,
    DataPartitionPlan,
    Dataset,
    DatasetPipelineRun,
    Project,
    ProjectMembership,
    Span,
    User,
)
from overbae.services.datasets import files, paths

sys.path.insert(0, str(Path.cwd() / "overmind"))
from overmind.dataset_cmd import (  # noqa: E402
    DatasetUploadError,
    export_dataset,
    upload_file,
    wait_until_ready,
)

BASE = os.environ.get("WORKSHOP_ACCEPTANCE_URL", "http://127.0.0.1:8000")
SIZES = [
    int(size)
    for size in os.environ.get("WORKSHOP_ACCEPTANCE_ROWS", "100000,1000000").split(",")
    if size.strip()
]
EXPECTED = {
    "start_dataset",
    "cancel_dataset",
    "list_datasets",
    "inspect_dataset",
    "query_dataset",
    "create_dataset_from_traces",
    "create_dataset_from_llm_calls",
    "inspect_dataset_workbench",
    "save_dataset_pipeline",
    "run_dataset_pipeline",
    "import_dataset_version",
    "update_dataset",
    "explore_dataset",
    "derive_dataset",
    "retry_data_partition",
    "create_data_partition",
    "get_job",
    "list_model_workflows",
    "list_projects",
}


class Replay:
    def __init__(self, directory):
        suffix = uuid.uuid4().hex
        self.directory = Path(directory)
        self.user = User.objects.create_user(email=f"workshop-acceptance-{suffix}@example.test")
        self.project = Project.objects.create(name="Workshop acceptance", slug=f"workshop-{suffix}")
        self.foreign = Project.objects.create(name="Foreign acceptance", slug=f"foreign-{suffix}")
        ProjectMembership.objects.create(user=self.user, project=self.project)
        self.key, _ = APIToken.create_for_user(self.user, project=self.project)
        self.read_key, _ = APIToken.create_for_user(
            self.user, project=self.project, permission=["read"]
        )
        self.account_key, _ = APIToken.create_for_user(self.user)
        self.http = httpx.Client(timeout=60)
        self.results = []
        self.called = set()
        self.success_counts = Counter()
        self.error_counts = Counter()
        self.resources = set()
        self.max_response = 0
        self.pdf_uploads = []
        self.rpc_measurements = []
        self.job_snapshots = {}
        self.unchanged_job_reads = 0

    def rpc(self, method, params=None, *, key=None):
        started = time.monotonic()
        response = self.http.post(
            BASE + "/api/mcp/",
            headers={"X-Api-Key": key or self.key, "Accept": "application/json"},
            json={"jsonrpc": "2.0", "id": 1, "method": method, "params": params or {}},
        )
        self.rpc_measurements.append(
            {
                "operation": (params or {}).get("name", method),
                "milliseconds": (time.monotonic() - started) * 1000,
                "bytes": len(response.content),
                "http_status": response.status_code,
            }
        )
        assert response.status_code == 200, (response.status_code, response.text[:500])
        self.max_response = max(self.max_response, len(response.content))
        value = response.json()
        assert "error" not in value, value
        return value["result"]

    def call(self, name, arguments=None, *, error=False, key=None):
        result = self.rpc("tools/call", {"name": name, "arguments": arguments or {}}, key=key)
        value = result.get("structuredContent")
        text = next(item["text"] for item in result["content"] if item["type"] == "text")
        assert json.loads(text) == value
        if error:
            assert result.get("isError"), (name, value)
            assert value["error"]["code"] != "internal_error", (name, value)
        else:
            assert not result.get("isError"), (name, value)
            self.called.add(name)
            self.success_counts[name] += 1
        if error:
            self.error_counts[name] += 1
        if name == "get_job" and not error:
            identity = (value["kind"], value["id"])
            snapshot = json.dumps(
                {field: value.get(field) for field in ("status", "progress", "job_error")},
                sort_keys=True,
            )
            if self.job_snapshots.get(identity) == snapshot:
                self.unchanged_job_reads += 1
            self.job_snapshots[identity] = snapshot
        return value

    def rpc_performance(self):
        groups = {}
        for measurement in self.rpc_measurements:
            groups.setdefault(measurement["operation"], []).append(measurement)
        result = {}
        for operation, measurements in sorted(groups.items()):
            times = sorted(item["milliseconds"] for item in measurements)
            count = len(times)
            result[operation] = {
                "calls": count,
                "p50_ms": round(times[(count - 1) // 2], 3),
                "p95_ms": round(times[(95 * count + 99) // 100 - 1], 3),
                "max_ms": round(times[-1], 3),
                "total_bytes": sum(item["bytes"] for item in measurements),
                "max_bytes": max(item["bytes"] for item in measurements),
                "non_200_responses": sum(item["http_status"] != 200 for item in measurements),
            }
        return {
            "basis": "Client-observed local HTTP round trips; includes expected validation rejections, excludes local parsing and asynchronous job duration.",
            "operations": result,
            "unchanged_job_reads": self.unchanged_job_reads,
            "distinct_jobs_read": len(self.job_snapshots),
        }

    def resource(self, uri, *, key=None):
        result = self.rpc("resources/read", {"uri": uri}, key=key)
        value = json.loads(result["contents"][0]["text"])
        self.resources.add(uri)
        return value

    def records(self, dataset, cell=None):
        result = []
        while True:
            page = self.call(
                "query_dataset",
                {
                    "dataset": dataset,
                    **({"cell": cell} if cell else {}),
                    "sql": f"SELECT * FROM t ORDER BY source_row LIMIT 100 OFFSET {len(result)}",
                },
            )
            result.extend(page["rows"])
            if len(page["rows"]) < 100:
                return result

    def simultaneous(self, name, arguments, count=8):
        barrier = Barrier(count)

        def submit(_):
            barrier.wait(timeout=20)
            return self.call(name, arguments)

        with ThreadPoolExecutor(max_workers=count) as pool:
            return list(pool.map(submit, range(count)))

    def check(self, label, function):
        selected = os.environ.get("WORKSHOP_ACCEPTANCE_CASES", "").split(",")
        if selected != [""] and label not in selected:
            return
        start = time.monotonic()
        measurement_start = len(self.rpc_measurements)
        try:
            details = function() or {}
            row = {
                "case": label,
                "passed": True,
                "seconds": round(time.monotonic() - start, 3),
                **details,
            }
        except Exception as exc:
            row = {
                "case": label,
                "passed": False,
                "seconds": round(time.monotonic() - start, 3),
                "error": str(exc)[:1200],
                "traceback": traceback.format_exc(limit=4),
            }
        measurements = self.rpc_measurements[measurement_start:]
        row["mcp_calls"] = len(measurements)
        row["mcp_response_bytes"] = sum(item["bytes"] for item in measurements)
        self.results.append(row)
        print(json.dumps(row, ensure_ascii=False), flush=True)

    def inspect(self, dataset):
        return self.call("inspect_dataset", {"dataset": dataset})

    def source_inventory(self, dataset):
        sources = []
        offset = 0
        while True:
            detail = self.call(
                "inspect_dataset", {"dataset": dataset, "source_offset": offset, "source_limit": 10}
            )
            assert len(json.dumps(detail, ensure_ascii=False).encode()) <= 32_000
            sources.extend(detail["sources"])
            page = detail["source_page"]
            assert page["offset"] == offset and page["total"] == detail["sources_total"]
            if not page["has_more"]:
                assert len(sources) == page["total"]
                return sources
            assert int(page["next_cursor"]) > offset
            offset = int(page["next_cursor"])

    def active(self, dataset):
        detail = self.inspect(dataset)
        return detail["active"]

    def poll(self, kind, identifier, *, expected="completed", timeout=600):
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            result = self.call("get_job", {"kind": kind, "id": identifier})
            if result["status"] in {"completed", "failed", "cancelled", "idle", "error", "ready"}:
                assert result["status"] == expected, result
                return result
            time.sleep(0.25)
        raise AssertionError(f"{kind} {identifier} exceeded {timeout}s")

    def upload(self, filename, records, *, dataset=None, intent="explore"):
        path = self.directory / filename
        with path.open("w") as output:
            for row in records:
                output.write(json.dumps(row, ensure_ascii=False) + "\n")
        result = upload_file(
            path,
            project_id=str(self.project.pk),
            api_key=self.key,
            api_url=BASE,
            intent=None if dataset else intent,
            dataset=dataset,
        )
        wait_until_ready(result["id"], api_key=self.key, api_url=BASE, poll=0.25, timeout=600)
        return result["id"]

    def pipeline(
        self, dataset, source, steps, *, name="Transform", request_key=None, expected="completed"
    ):
        key = request_key or uuid.uuid4().hex
        recipe = self.call(
            "save_dataset_pipeline",
            {"dataset": dataset, "name": name, "request_key": key, "steps": steps},
        )["pipeline"]
        arguments = {
            "dataset": dataset,
            "pipeline": recipe["id"],
            "source_cell": source["id"],
            "source_fingerprint": source["fingerprint"],
            "request_key": key,
        }
        run = self.call("run_dataset_pipeline", arguments)
        result = self.poll("dataset_pipeline", run["job"]["id"], expected=expected)
        if expected == "completed":
            assert result["progress"]["stage"] == "completed"
            assert result["progress"]["source_rows"] == source["rows"]
        return recipe, arguments, result

    def discovery(self):
        initialized = self.rpc(
            "initialize",
            {
                "protocolVersion": "2025-03-26",
                "capabilities": {},
                "clientInfo": {"name": "native-workshop-acceptance", "version": "1"},
            },
        )
        assert initialized["serverInfo"]["version"] == "4.0.0", initialized
        tools = self.rpc("tools/list")["tools"]
        names = {tool["name"] for tool in tools}
        assert names >= EXPECTED
        assert not names & {"message_dataset_agent", "run_dataset", "manage_dataset_workflow"}
        workshop = [tool for tool in tools if tool["name"] in EXPECTED]
        for tool in workshop:
            Draft202012Validator.check_schema(tool["inputSchema"])
            self.call(tool["name"], {"invalid_extra_field": True}, error=True)
        self.rpc("resources/read", {"uri": "overmind://interface/current"})
        assert self.rpc("prompts/list")["prompts"]
        projects = self.call("list_projects")
        assert [item["id"] for item in projects["projects"]] == [str(self.project.pk)]
        readonly = {tool["name"] for tool in self.rpc("tools/list", key=self.read_key)["tools"]}
        assert "inspect_dataset" in readonly and "import_dataset_version" not in readonly
        self.call("start_dataset", {"brief": "Do not create"}, error=True, key=self.read_key)
        self.call("list_datasets", {}, error=True, key=self.account_key)
        self.call("list_datasets", {"project_id": str(self.project.pk)}, key=self.account_key)
        self.call(
            "list_datasets", {"project_id": str(self.foreign.pk)}, error=True, key=self.account_key
        )
        return {"catalogue_tools": len(tools), "workshop_argument_checks": len(workshop)}

    def draft_and_semantic_import(self):
        draft = self.call(
            "start_dataset", {"name": "Library", "brief": "Create grounded training examples."}
        )
        dataset = draft["dataset"]["id"]
        self.call(
            "update_dataset", {"dataset": dataset, "intent": "train", "name": "Library examples"}
        )
        self.upload(
            "library.jsonl",
            [{"text": "The library opens at nine. Members may borrow four books."}],
            dataset=dataset,
        )
        source = self.active(dataset)
        assert self.inspect(dataset)["brief"] == "Create grounded training examples."
        records = [
            {
                "messages": [
                    {
                        "role": "user",
                        "content": "Passage: The library opens at nine. Members may borrow four books.\nWhen does the library open?",
                    },
                    {"role": "assistant", "content": "At nine."},
                ],
                "source_row": 0,
                "human_reviewed": True,
            },
            {
                "messages": [
                    {
                        "role": "user",
                        "content": "Passage: The library opens at nine. Members may borrow four books.\nHow many books may a member borrow?",
                    },
                    {"role": "assistant", "content": "Four books."},
                ],
                "_overmind_parent_rows": [0],
            },
        ]
        args = {
            "dataset": dataset,
            "source_cell": source["id"],
            "source_fingerprint": source["fingerprint"],
            "request_key": "native-authored",
            "name": "Grounded examples",
            "provenance": "Authored by the native coding agent from the pinned library passage.",
            "imported_rows": records,
        }
        run = self.call("import_dataset_version", args)
        self.poll("dataset_pipeline", run["job"]["id"])
        repeated = self.call("import_dataset_version", args)
        assert repeated["run"]["id"] == run["run"]["id"]
        rows = self.call(
            "query_dataset", {"dataset": dataset, "sql": "SELECT * FROM t ORDER BY source_row"}
        )["rows"]
        assert len(rows) == 2 and all("human_reviewed" not in row for row in rows)
        self.call("import_dataset_version", {**args, "name": "Changed"}, error=True)
        self.call(
            "import_dataset_version",
            {**args, "request_key": "stale", "source_fingerprint": "0" * 64},
            error=True,
        )
        self.rpc("resources/read", {"uri": run["job"]["resource"]["uri"]})
        self.call("cancel_dataset", {"dataset": dataset})
        self.small = dataset
        self.small_source = source
        return {"published_rows": 2}

    def boundaries(self):
        dataset, source = self.small, self.small_source
        for sql in (
            "DELETE FROM t",
            "SELECT * FROM read_csv_auto('/etc/passwd')",
            "SELECT 1; SELECT 2",
        ):
            self.call("query_dataset", {"dataset": dataset, "sql": sql}, error=True)
        before = self.active(dataset)["id"]
        for steps in (
            [{"operation": "python", "code": "pass"}],
            [{"operation": "select", "columns": ["source_row"]}],
            [{"operation": "rename", "mapping": {"a": "x", "b": "x"}}],
        ):
            self.call(
                "save_dataset_pipeline",
                {
                    "dataset": dataset,
                    "name": "Invalid",
                    "request_key": uuid.uuid4().hex,
                    "steps": steps,
                },
                error=True,
            )
        for steps in (
            [{"operation": "select", "columns": ["missing"]}],
            [{"operation": "filter", "column": "text", "equals": "absent"}],
        ):
            self.pipeline(dataset, source, steps, expected="failed")
            assert self.active(dataset)["id"] == before
        base = {
            "dataset": dataset,
            "source_cell": source["id"],
            "source_fingerprint": source["fingerprint"],
            "request_key": "limit",
            "name": "Limits",
            "provenance": "Acceptance fixture",
        }
        for records in (
            [],
            [{"text": "row", "source_row": 0}] * 2001,
        ):
            self.call("import_dataset_version", {**base, "imported_rows": records}, error=True)
        oversized = self.http.post(
            BASE + "/api/mcp/",
            headers={"X-Api-Key": self.key},
            json={
                "jsonrpc": "2.0",
                "id": 1,
                "method": "tools/call",
                "params": {
                    "name": "import_dataset_version",
                    "arguments": {
                        **base,
                        "imported_rows": [{"text": "x" * (4 * 1024 * 1024), "source_row": 0}],
                    },
                },
            },
        )
        assert oversized.status_code == 413
        assert oversized.json()["error"]["code"] == "invalid_request"
        assert self.active(dataset)["id"] == before
        for records in ([{"text": "unbound"}], [{"text": "wrong parent", "source_row": 999999}]):
            run = self.call(
                "import_dataset_version",
                {**base, "request_key": uuid.uuid4().hex, "imported_rows": records},
            )
            self.poll("dataset_pipeline", run["job"]["id"], expected="failed")
            assert self.active(dataset)["id"] == before
        records = [
            {
                "text": f"value {i} — 日本語 🧪",
                "source_row": 0,
                "nullable": None,
                "nested": {"list": [None, True, 0, "é"]},
            }
            for i in range(2000)
        ]
        run = self.call("import_dataset_version", {**base, "imported_rows": records})
        self.poll("dataset_pipeline", run["job"]["id"])
        assert self.active(dataset)["rows"] == 2000
        self.call("update_dataset", {"dataset": dataset, "active": before})
        foreign = Dataset.objects.create(
            project=self.foreign, name="Foreign source", state=Dataset.State.IDLE
        )
        for name in (
            "inspect_dataset",
            "inspect_dataset_workbench",
            "cancel_dataset",
            "update_dataset",
        ):
            self.call(
                name,
                {
                    "dataset": str(foreign.pk),
                    **({"name": "Forbidden"} if name == "update_dataset" else {}),
                },
                error=True,
            )

    def large(self, count):
        stages = {}
        stage_started = time.monotonic()
        print(json.dumps({"progress": "upload", "source_rows": count}), flush=True)
        dataset = self.upload(
            f"large-{count}.jsonl",
            (
                {
                    "question": f"What is item {i}?",
                    "answer": f"Item {i} — café",
                    "segment": i % 2,
                    "value": i,
                    "group_id": f"group-{i // 2}",
                }
                for i in range(count)
            ),
            intent="train",
        )
        source = self.active(dataset)
        assert source["rows"] == count
        stages["source_creation_upload_landing"] = round(time.monotonic() - stage_started, 3)
        stage_started = time.monotonic()
        print(json.dumps({"progress": "transform", "source_rows": count}), flush=True)
        recipe, args, result = self.pipeline(
            dataset,
            source,
            [
                {"operation": "rename", "mapping": {"question": "prompt"}},
                {"operation": "filter", "column": "segment", "equals": 0},
                {"operation": "conversation", "question": "prompt", "answer": "answer"},
                {"operation": "select", "columns": ["messages", "value", "group_id"]},
            ],
            name="Even conversations",
            request_key=f"large-{count}",
        )
        output = self.active(dataset)
        stages["deterministic_pipeline"] = round(time.monotonic() - stage_started, 3)
        expected = (count + 1) // 2
        assert output["rows"] == expected
        with ThreadPoolExecutor(max_workers=4) as pool:
            receipts = list(pool.map(lambda _: self.call("run_dataset_pipeline", args), range(8)))
        assert {receipt["run"]["id"] for receipt in receipts} == {result["id"]}
        values = self.call(
            "query_dataset",
            {"dataset": dataset, "sql": "SELECT count(*) AS n, sum(value) AS total FROM t"},
        )["rows"][0]
        assert values["n"] == expected and int(values["total"]) == expected * (expected - 1), values
        exported = self.directory / f"export-{count}.jsonl"
        stage_started = time.monotonic()
        export_dataset(dataset, cell=output["id"], output=exported, api_key=self.key, api_url=BASE)
        stages["export"] = round(time.monotonic() - stage_started, 3)
        stage_started = time.monotonic()
        derived = self.directory / f"derived-{count}.jsonl"
        with exported.open() as incoming, derived.open("w") as outgoing:
            for line in incoming:
                row = json.loads(line)
                row["_overmind_parent_rows"] = [row["source_row"]]
                row["derived_value"] = row["value"] * 3
                outgoing.write(json.dumps(row, ensure_ascii=False) + "\n")
        stages["local_derivation"] = round(time.monotonic() - stage_started, 3)
        stage_started = time.monotonic()
        print(json.dumps({"progress": "upload_external_artifact", "rows": expected}), flush=True)
        artifact = upload_file(
            derived, project_id=str(self.project.pk), api_key=self.key, api_url=BASE, intent="train"
        )
        wait_until_ready(artifact["id"], api_key=self.key, api_url=BASE, poll=0.25, timeout=600)
        artifact_cell = self.active(artifact["id"])
        stages["artifact_upload_landing"] = round(time.monotonic() - stage_started, 3)
        stage_started = time.monotonic()
        imported = self.call(
            "import_dataset_version",
            {
                "dataset": dataset,
                "source_cell": output["id"],
                "source_fingerprint": output["fingerprint"],
                "artifact_cell": artifact_cell["id"],
                "artifact_fingerprint": artifact_cell["fingerprint"],
                "name": "Local derivation",
                "provenance": "Streamed native-agent-authored multiplication over the exported pinned version.",
                "request_key": f"artifact-{count}",
            },
        )
        print(json.dumps({"progress": "import_external_artifact", "rows": expected}), flush=True)
        imported_job = self.poll("dataset_pipeline", imported["job"]["id"])
        stages["lineage_bound_import"] = round(time.monotonic() - stage_started, 3)
        values = self.call(
            "query_dataset",
            {"dataset": dataset, "sql": "SELECT count(*) AS n, sum(derived_value) AS total FROM t"},
        )["rows"][0]
        assert values["n"] == expected and int(values["total"]) == 3 * expected * (expected - 1), (
            values
        )
        original = self.call(
            "query_dataset",
            {"dataset": dataset, "cell": source["id"], "sql": "SELECT count(*) AS n FROM t"},
        )
        assert original["rows"][0]["n"] == count
        self.call("inspect_dataset_workbench", {"dataset": dataset})
        self.large_dataset = dataset
        return {
            "source_rows": count,
            "transformed_rows": expected,
            "imported_rows": expected,
            "export_bytes": exported.stat().st_size,
            "duplicate_retries": 8,
            "stage_seconds": stages,
            "pipeline_stage_seconds": result["progress"].get("stage_seconds"),
            "import_stage_seconds": imported_job["progress"].get("stage_seconds"),
        }

    def explore_partition(self):
        dataset = self.upload(
            "partitions.jsonl",
            (
                {
                    "messages": [
                        {"role": "user", "content": f"Question {i}"},
                        {"role": "assistant", "content": f"Answer {i}"},
                    ],
                    "group_id": f"group-{i // 2}",
                }
                for i in range(1000)
            ),
            intent="train",
        )
        source = self.active(dataset)
        for tool in ("explore_dataset", "derive_dataset"):
            op = self.call(tool, {"source_cell": source["id"], "name": tool, "request_key": tool})
            self.poll("data_exploration", op["workflow"]["id"])
        plan = self.call(
            "create_data_partition",
            {
                "source_cell": source["id"],
                "name": "Grouped split",
                "request_key": "split",
                "recipe": {
                    "seed": 17,
                    "fractions": {"train": 0.8, "final": 0.2},
                    "group_by": ["group_id"],
                },
            },
        )
        self.poll("data_partition", plan["workflow"]["id"])
        self.call("retry_data_partition", {"partition": plan["workflow"]["id"]}, error=True)
        DataPartitionPlan.objects.filter(pk=plan["workflow"]["id"]).update(
            state="failed", error="Injected acceptance interruption"
        )
        retry = self.call("retry_data_partition", {"partition": plan["workflow"]["id"]})
        assert retry["workflow"]["id"] == plan["workflow"]["id"]
        self.poll("data_partition", plan["workflow"]["id"])
        self.call("list_model_workflows", {"kind": "data_partition"})
        self.call("update_dataset", {"dataset": dataset, "intent": "eval"}, error=True)

    def typed_and_stepwise(self):
        decisions = [
            {
                "state": "",
                "question": "Choose",
                "kind": "choice",
                "options": ["B", "A"],
                "target_probabilities": [0.3, 0.7],
                "weight": 2.5,
            },
            {
                "state": "",
                "question": "Choose",
                "kind": "choice",
                "options": ["B", "A"],
                "target_probabilities": [0.3, 0.7],
                "weight": 2.5,
            },
            {
                "state": "Café 🧪",
                "question": "Rate",
                "kind": "score",
                "options": ["low", "middle", "high"],
                "target_probabilities": [0.1, 0.2, 0.7],
                "weight": 0.25,
            },
        ]
        dataset = self.upload(
            "typed-decisions.jsonl",
            (
                {"decision": decision, "group_id": f"group-{i // 2}"}
                for i, decision in enumerate(decisions)
            ),
            intent="train",
        )
        source = self.active(dataset)
        self.pipeline(
            dataset, source, [{"operation": "select", "columns": ["decision", "group_id"]}]
        )
        selected = self.active(dataset)
        records = self.call(
            "query_dataset", {"dataset": dataset, "sql": "SELECT * FROM t ORDER BY source_row"}
        )["rows"]
        assert [row["decision"] for row in records] == decisions
        receipt = self.call(
            "import_dataset_version",
            {
                "dataset": dataset,
                "source_cell": selected["id"],
                "source_fingerprint": selected["fingerprint"],
                "name": "Preserved native targets",
                "provenance": "Identity copy by the acceptance native client.",
                "request_key": "typed",
                "imported_rows": records,
            },
        )
        self.poll("dataset_pipeline", receipt["job"]["id"])
        preserved = self.call(
            "query_dataset", {"dataset": dataset, "sql": "SELECT * FROM t ORDER BY source_row"}
        )["rows"]
        assert [row["decision"] for row in preserved] == decisions
        assert [row["group_id"] for row in preserved] == ["group-0", "group-0", "group-1"]

        chain = self.upload(
            "stepwise.jsonl",
            ({"question": f"Q{i}", "answer": f"A{i}", "keep": i % 2} for i in range(8)),
            intent="train",
        )
        sources = [self.active(chain)]
        steps = [
            {"operation": "rename", "mapping": {"question": "prompt"}},
            {"operation": "filter", "column": "keep", "equals": 0},
            {"operation": "conversation", "question": "prompt", "answer": "answer"},
            {"operation": "select", "columns": ["messages"]},
        ]
        for index, step in enumerate(steps):
            self.pipeline(chain, sources[-1], [step], name=f"Step {index + 1}")
            sources.append(self.active(chain))
        assert len({cell["id"] for cell in sources}) == 5
        for cell, expected in zip(sources, [8, 8, 4, 4, 4], strict=True):
            result = self.call(
                "query_dataset",
                {"dataset": chain, "cell": cell["id"], "sql": "SELECT count(*) AS n FROM t"},
            )
            assert result["rows"][0]["n"] == expected
        self.upload(
            "later-source.jsonl",
            [
                {
                    "messages": [
                        {"role": "user", "content": "Later question"},
                        {"role": "assistant", "content": "Later answer"},
                    ]
                }
            ],
            dataset=chain,
        )
        assert self.active(chain)["rows"] == 5
        assert (
            self.call(
                "query_dataset",
                {"dataset": chain, "cell": sources[-1]["id"], "sql": "SELECT count(*) AS n FROM t"},
            )["rows"][0]["n"]
            == 4
        )
        return {"typed_decisions": 3, "stepwise_versions": 5, "appended_source_rows": 1}

    def trace_sources(self):
        capability = Capability(
            id=uuid.uuid4(), project=self.project, name="Fixture", slug="fixture"
        )
        spans = []
        for index in range(20):
            spans.append(
                Span(
                    span_id=uuid.uuid4().hex[:16],
                    trace_id=uuid.uuid4().hex,
                    project=self.project,
                    capability=capability,
                    span_type="llm_call",
                    name="Recorded fixture",
                    start_time_ns=1,
                    end_time_ns=2,
                    attributes={
                        "overmind.input.data": json.dumps(
                            [{"role": "user", "content": f"Question {index}"}]
                        ),
                        "overmind.output.data": json.dumps(
                            [{"role": "assistant", "content": f"Answer {index}"}]
                        ),
                        "genai.model": "fixture/local",
                    },
                )
            )
        # Backdate before commit so the live scoring sweep cannot claim these synthetic spans.
        with transaction.atomic():
            Capability.objects.bulk_create([capability])
            Span.objects.bulk_create(spans)
            Span.objects.filter(project=self.project).update(
                received_at=timezone.now() - timedelta(days=1)
            )
        traces = self.call(
            "create_dataset_from_traces",
            {
                "name": "Trace source",
                "trace_ids": [span.trace_id for span in spans],
                "capability": None,
            },
        )
        self.poll("dataset_run", traces["dataset"]["id"], expected="idle")
        assert self.active(traces["dataset"]["id"])["rows"] == 20
        calls = self.call(
            "create_dataset_from_llm_calls",
            {
                "name": "Call source",
                "capability": str(capability.pk),
                "since": (timezone.now() - timedelta(days=2)).isoformat(),
                "intent": "train",
            },
        )
        self.poll("dataset_run", calls["dataset"]["id"], expected="idle")
        assert self.active(calls["dataset"]["id"])["rows"] == 20
        for tool, arguments in (
            (
                "create_dataset_from_traces",
                {
                    "name": "Trace split",
                    "trace_ids": [span.trace_id for span in spans],
                    "split": {"eval_percent": 25},
                    "capability": None,
                },
            ),
            (
                "create_dataset_from_llm_calls",
                {
                    "name": "Call split",
                    "capability": str(capability.pk),
                    "since": (timezone.now() - timedelta(days=2)).isoformat(),
                    "split": True,
                    "eval_percent": 25,
                },
            ),
        ):
            split = self.call(tool, arguments)
            parts = []
            for field, intent in (("dataset", "train"), ("eval_dataset", "eval")):
                identifier = split[field]["id"]
                self.poll("dataset_run", identifier, expected="idle")
                detail = self.inspect(identifier)
                assert detail["intent"] == intent
                parts.append(self.records(identifier))
            assert sum(len(part) for part in parts) == 20
            # Trace identity, not freshly assigned row positions, defines cross-partition overlap.
            identities = [
                {
                    row.get("trace_id") or row.get("source_trace_id") or row.get("origin_trace_id")
                    for row in part
                }
                for part in parts
            ]
            assert None not in set.union(*identities), parts
            assert not identities[0] & identities[1]
            self.call(tool, {**arguments, "name": "x" * 255}, error=True)
        filtered = self.call(
            "create_dataset_from_traces",
            {
                "name": "Selected traces",
                "filters": {"capability": str(capability.pk)},
                "limit": 7,
                "capability": None,
            },
        )
        self.poll("dataset_run", filtered["dataset"]["id"], expected="idle")
        assert self.active(filtered["dataset"]["id"])["rows"] == 7
        for arguments in (
            {"trace_ids": ["f" * 32]},
            {"filters": {"invented_filter": True}},
        ):
            self.call("create_dataset_from_traces", {"name": "Rejected", **arguments}, error=True)
        self.call(
            "create_dataset_from_llm_calls",
            {"name": "No calls", "capability": str(capability.pk), "since": "2999-01-01T00:00:00Z"},
            error=True,
        )
        for since in ("not-a-date", "2026-99-01T00:00:00Z"):
            self.call(
                "create_dataset_from_llm_calls",
                {"name": "Invalid window", "capability": str(capability.pk), "since": since},
                error=True,
            )
        naive = self.call(
            "create_dataset_from_llm_calls",
            {
                "name": "Naive UTC window",
                "capability": str(capability.pk),
                "since": "2020-01-01T00:00:00",
                "until": "2999-01-01T00:00:00",
                "limit": 3,
            },
        )
        self.poll("dataset_run", naive["dataset"]["id"], expected="idle")
        assert self.active(naive["dataset"]["id"])["rows"] == 3
        return {"trace_rows": 20, "call_rows": 20, "split_observations": 40, "filtered_rows": 7}

    def operation_matrix(self):
        records = [
            {
                "left": f"Q{i}",
                "right": f"A{i}",
                "nullable": None if i % 2 else "present",
                "object": {"list": [i % 2, None, True, "日本語"]},
                "flag": bool(i % 2),
                'quote" café': "x" * 12000 if i == 0 else "line\n二",
            }
            for i in range(6)
        ]
        dataset = self.upload("operation-matrix.jsonl", records)
        source = self.active(dataset)
        cases = [
            ([{"operation": "rename", "mapping": {"left": "right", "right": "left"}}], 6),
            ([{"operation": "filter", "column": "nullable", "equals": None}], 3),
            ([{"operation": "filter", "column": "flag", "equals": True}], 3),
            ([{"operation": "filter", "column": "object", "equals": records[0]["object"]}], 3),
            ([{"operation": "select", "columns": ['quote" café', "object"]}], 6),
            ([{"operation": "conversation", "question": "left", "answer": "right"}], 6),
        ]
        for steps, count in cases:
            self.pipeline(dataset, source, steps)
            output = self.records(dataset)
            assert len(output) == count
            for row in output:
                original = records[row["source_row"]]
                op = steps[0]["operation"]
                if op == "rename":
                    assert (row["left"], row["right"]) == (original["right"], original["left"])
                elif op == "conversation":
                    assert row["messages"] == [
                        {"role": "user", "content": original["left"]},
                        {"role": "assistant", "content": original["right"]},
                    ]
                else:
                    for key, value in original.items():
                        if key in row:
                            assert row[key] == value, (op, key)
        before = self.active(dataset)
        for steps in (
            [{"operation": "rename", "mapping": {"left": "nullable"}}],
            [{"operation": "conversation", "question": "left", "answer": "object"}],
            [{"operation": "conversation", "question": "left", "answer": "nullable"}],
        ):
            self.pipeline(dataset, source, steps, expected="failed")
            assert self.active(dataset) == before
        self.pipeline(
            dataset,
            before,
            [{"operation": "conversation", "question": "left", "answer": "right"}],
            expected="failed",
        )
        assert self.active(dataset) == before
        original = self.records(dataset, source["id"])
        for row, expected in zip(original, records, strict=True):
            assert all(row[key] == value for key, value in expected.items())
        return {
            "successful_operations": len(cases),
            "atomic_failures": 4,
            "long_value_chars": 12000,
        }

    def concurrent_first_submissions(self):
        dataset = self.upload("concurrent.jsonl", [{"text": f"row {i}"} for i in range(32)])
        source = self.active(dataset)
        save = {
            "dataset": dataset,
            "name": "Select",
            "request_key": "concurrent-recipe",
            "steps": [{"operation": "select", "columns": ["text"]}],
        }
        recipes = self.simultaneous("save_dataset_pipeline", save)
        recipe_ids = {result["pipeline"]["id"] for result in recipes}
        assert len(recipe_ids) == 1
        self.call("save_dataset_pipeline", {**save, "name": "Changed"}, error=True)
        run_args = {
            "dataset": dataset,
            "pipeline": recipe_ids.pop(),
            "source_cell": source["id"],
            "source_fingerprint": source["fingerprint"],
            "request_key": "concurrent-run",
        }
        runs = self.simultaneous("run_dataset_pipeline", run_args)
        assert len({result["run"]["id"] for result in runs}) == 1
        self.poll("dataset_pipeline", runs[0]["run"]["id"])
        self.call("run_dataset_pipeline", {**run_args, "source_fingerprint": "0" * 64}, error=True)
        imports = self.simultaneous(
            "import_dataset_version",
            {
                "dataset": dataset,
                "source_cell": source["id"],
                "source_fingerprint": source["fingerprint"],
                "request_key": "concurrent-import",
                "name": "Copied",
                "provenance": "Concurrent native client identity copy",
                "imported_rows": [{"text": f"row {i}", "source_row": i} for i in range(32)],
            },
        )
        assert len({result["run"]["id"] for result in imports}) == 1
        self.poll("dataset_pipeline", imports[0]["run"]["id"])
        detail = self.call("inspect_dataset", {"dataset": dataset, "cell_limit": 20})
        assert len(detail["cells"]) == 3
        assert len(self.call("inspect_dataset_workbench", {"dataset": dataset})["runs"]) == 2
        return {"simultaneous_first_requests": 24, "published_versions": 2, "recipes": 1}

    def all_parent_lineage(self):
        dataset = self.upload(
            "parents.jsonl", [{"text": f"Evidence {i}", "group_id": f"g{i // 2}"} for i in range(4)]
        )
        source = self.active(dataset)
        base = {
            "dataset": dataset,
            "source_cell": source["id"],
            "source_fingerprint": source["fingerprint"],
            "name": "Merged and split",
            "provenance": "Native fixture combines exact parent text",
        }
        for parents in ([], [999], [0, 999], [True], ["0"], [-1], [0.5], [{"row": 0}]):
            receipt = self.call(
                "import_dataset_version",
                {
                    **base,
                    "request_key": uuid.uuid4().hex,
                    "imported_rows": [
                        {"text": "valid first", "source_row": 0},
                        {"text": "invalid second", "_overmind_parent_rows": parents},
                    ],
                },
            )
            self.poll("dataset_pipeline", receipt["run"]["id"], expected="failed")
            assert self.active(dataset) == source
        receipt = self.call(
            "import_dataset_version",
            {
                **base,
                "request_key": "merge",
                "imported_rows": [
                    {
                        "text": "Evidence 0 + Evidence 1",
                        "_overmind_parent_rows": [0, 1, 0],
                        "_overmind_provenance": {"verified": True},
                        "human_reviewed": True,
                    },
                    {"text": "Evidence 2 first part", "source_row": 2},
                    {"text": "Evidence 2 second part", "source_row": 2},
                ],
            },
        )
        self.poll("dataset_pipeline", receipt["run"]["id"])
        output = self.records(dataset)
        assert len({row["source_row"] for row in output}) == 3
        by_text = {row["text"]: row for row in output}
        merged = by_text["Evidence 0 + Evidence 1"]
        assert {parent["row"] for parent in merged["_overmind_provenance"]["parents"]} == {0, 1}
        for row in output:
            assert "human_reviewed" not in row
            assert "verified" not in row["_overmind_provenance"]
            assert all(p["cell"] == source["id"] for p in row["_overmind_provenance"]["parents"])
        self.pipeline(dataset, self.active(dataset), [{"operation": "select", "columns": ["text"]}])
        after_projection = self.records(dataset)
        assert [row["_overmind_provenance"] for row in after_projection] == [
            row["_overmind_provenance"] for row in output
        ]
        return {"invalid_parent_shapes": 8, "preserved_parents": [0, 1, 2], "published_rows": 3}

    def navigation_and_query(self):
        for uri in (
            "overmind://interface/current",
            "overmind://dataset-upload",
            "overmind://dataset-export",
        ):
            assert self.resource(uri)
        for name, arguments in (
            ("upload-dataset-file", {"path": "/tmp/fixture.jsonl"}),
            ("export-dataset", {"dataset": self.small}),
        ):
            prompt = self.rpc("prompts/get", {"name": name, "arguments": arguments})
            assert prompt["messages"]
            assert not any(
                word in json.dumps(prompt)
                for word in ("message_dataset_agent", "manage_dataset_workflow")
            )
        dataset = self.upload("query-pages.jsonl", [{"value": i} for i in range(205)])
        source = self.active(dataset)
        first = self.call(
            "query_dataset", {"dataset": dataset, "sql": "SELECT * FROM t ORDER BY value"}
        )
        assert first["n"] == 100 and first["truncated"]
        assert [row["value"] for row in self.records(dataset)] == list(range(205))
        for reference in (source["id"], source["version"]):
            result = self.call(
                "query_dataset",
                {
                    "dataset": dataset,
                    "cell": reference,
                    "sql": "WITH selected AS (SELECT * FROM t WHERE value >= 200) SELECT sum(value) AS n FROM selected",
                },
            )
            assert len(result["rows"]) == 1 and int(result["rows"][0]["n"]) == 1010, result
        for sql in (
            "SELECT * FROM t JOIN read_json_auto('/etc/passwd') ON true",
            "COPY t TO '/tmp/workshop-forbidden.csv'",
            "INSTALL httpfs",
            "ATTACH '/tmp/database' AS other",
        ):
            self.call("query_dataset", {"dataset": dataset, "sql": sql}, error=True)
        assert self.resource(f"overmind://datasets/{dataset}")["active"]["id"] == source["id"]
        listed = []
        offset = 0
        while True:
            page = self.call("list_datasets", {"limit": 3, "offset": offset})
            listed += [item["id"] for item in page["datasets"]]
            if not page["page"]["has_more"]:
                break
            offset = int(page["page"]["next_cursor"])
        assert (
            len(listed) == len(set(listed)) == Dataset.objects.filter(project=self.project).count()
        )
        return {"query_rows_reconstructed": 205, "listed_datasets": len(listed), "prompts": 2}

    def exploration_results(self):
        dataset = self.upload(
            "strata.jsonl", [{"value": i, "stratum": i % 125} for i in range(1000)]
        )
        source = self.active(dataset)
        args = {
            "source_cell": source["id"],
            "name": "Census",
            "request_key": "census",
            "sampling": {"rows": 250, "seed": 42, "stratify_by": ["stratum"]},
        }
        result = self.call("explore_dataset", args)
        identifier = result["workflow"]["id"]
        self.poll("data_exploration", identifier)
        detail = self.resource(f"overmind://jobs/data_exploration/{identifier}")
        sampling = detail["report"]["sampling"]
        assert (
            sampling["source_rows"] == 1000 and sampling["strata"] == 125 and sampling["feasible"]
        )
        allocations = []
        for offset in (0, 100):
            page = self.resource(
                f"overmind://jobs/data_exploration/{identifier}/strata?offset={offset}&limit=100"
            )
            assert page["total"] == 125
            allocations += page["strata"]
        assert len(allocations) == len({json.dumps(row["values"]) for row in allocations}) == 125
        assert sum(row["count"] for row in allocations) == 1000
        assert sum(row["allocation"] for row in allocations) == 250
        assert self.call("explore_dataset", args)["workflow"]["id"] == identifier
        self.call("explore_dataset", {**args, "name": "Changed"}, error=True)
        cached = self.call("explore_dataset", {**args, "request_key": "cached"})["workflow"]["id"]
        self.poll("data_exploration", cached)
        assert (
            self.resource(f"overmind://jobs/data_exploration/{cached}")["report"]["reused_from"]
            == identifier
        )
        infeasible = self.call(
            "explore_dataset",
            {**args, "request_key": "infeasible", "sampling": {**args["sampling"], "rows": 100}},
        )["workflow"]["id"]
        self.poll("data_exploration", infeasible)
        assert not self.resource(f"overmind://jobs/data_exploration/{infeasible}")["report"][
            "sampling"
        ]["feasible"]
        derived = self.call(
            "derive_dataset",
            {"source_cell": source["id"], "name": "Independent chain", "request_key": "derived"},
        )
        self.poll("data_exploration", derived["workflow"]["id"])
        output = self.resource(f"overmind://jobs/data_exploration/{derived['workflow']['id']}")[
            "output_dataset"
        ]
        assert self.active(output)["rows"] == 1000
        self.pipeline(
            output, self.active(output), [{"operation": "filter", "column": "stratum", "equals": 0}]
        )
        assert self.active(output)["rows"] == 8
        assert self.active(dataset) == source
        assert self.inspect(dataset)["cells"][0]["frozen"]
        return {
            "whole_source_rows": 1000,
            "strata": 125,
            "derived_rows": 1000,
            "derived_filtered_rows": 8,
        }

    def partition_contents_and_retry(self):
        records = [
            {
                "input": f"Question {i // 2}",
                "expected_output": f"Answer {i // 2}",
                "group_id": f"group-{i // 4}",
                "holdout": i < 8,
                "stratum": (i // 4) % 2,
            }
            for i in range(200)
        ]
        dataset = self.upload("partition-evidence.jsonl", records, intent="train")
        source = self.active(dataset)
        args = {
            "source_cell": source["id"],
            "name": "All four roles",
            "request_key": "four-roles",
            "recipe": {
                "seed": 104,
                "fractions": {"train": 0.5, "development": 0.2, "calibration": 0.1, "final": 0.2},
                "group_by": ["group_id"],
                "stratify_by": "stratum",
                "holdouts": [{"field": "holdout", "values": [True], "role": "final"}],
            },
        }
        result = self.call("create_data_partition", args)
        identifier = result["workflow"]["id"]
        self.poll("data_partition", identifier)
        detail = self.resource(f"overmind://jobs/data_partition/{identifier}")
        roles = {}
        all_rows = []
        for member in detail["members"]:
            cell = member["cell"]
            rows = self.records(cell["dataset_id"], cell["id"])
            assert rows
            roles[member["role"]] = {row["group_id"] for row in rows}
            for row in rows:
                if row["holdout"]:
                    assert member["role"] == "final"
            all_rows += [(row["input"], row["expected_output"], row["group_id"]) for row in rows]
            assert self.inspect(cell["dataset_id"])["intent"] == (
                "eval" if member["role"] in {"calibration", "final"} else "train"
            )
        assert Counter(all_rows) == Counter(
            (r["input"], r["expected_output"], r["group_id"]) for r in records
        )
        assert set(roles) == {"train", "development", "calibration", "final"}
        for left in roles:
            for right in roles:
                if left != right:
                    assert not roles[left] & roles[right]
        assert self.call("create_data_partition", args)["workflow"]["id"] == identifier
        self.call("create_data_partition", {**args, "name": "Changed"}, error=True)
        for recipe in (
            {**args["recipe"], "fractions": {"train": 0.8, "final": 0.3}},
            {**args["recipe"], "fractions": {"train": -0.2, "final": 1.2}},
        ):
            self.call(
                "create_data_partition",
                {**args, "request_key": uuid.uuid4().hex, "recipe": recipe},
                error=True,
            )
        # Only the disposable plan's terminal state is fault-injected; retry uses the live worker.
        DataPartitionPlan.objects.filter(pk=identifier).update(
            state="failed", error="Acceptance interruption after member publication"
        )
        self.call("retry_data_partition", {"partition": identifier})
        self.poll("data_partition", identifier)
        repeated = self.resource(f"overmind://jobs/data_partition/{identifier}")
        assert repeated["members"] == detail["members"]
        assert repeated["report"]["assignments_sha256"] == detail["report"]["assignments_sha256"]
        return {
            "input_observations": 200,
            "output_observations": len(all_rows),
            "disjoint_roles": 4,
            "retry_members_unchanged": True,
        }

    def full_history(self):
        dataset = self.upload("history.jsonl", [{"text": "original evidence"}])
        source = self.active(dataset)
        recipes, runs = set(), set()
        for index in range(105):
            recipe, _, result = self.pipeline(
                dataset,
                source,
                [{"operation": "select", "columns": ["text"]}],
                name=f"History {index}",
                request_key=f"history-{index}",
            )
            recipes.add(recipe["id"])
            runs.add(result["id"])
        found_recipes, found_runs, cells = [], [], []
        for offset in range(0, 120, 20):
            workbench = self.call(
                "inspect_dataset_workbench",
                {"dataset": dataset, "pipeline_offset": offset, "run_offset": offset, "limit": 20},
            )
            assert workbench["pipeline_page"]["total"] == workbench["run_page"]["total"] == 105
            found_recipes.extend(row["id"] for row in workbench["pipelines"])
            found_runs.extend(row["id"] for row in workbench["runs"])
        offset = 0
        while True:
            page = self.call(
                "inspect_dataset", {"dataset": dataset, "cell_offset": offset, "cell_limit": 20}
            )
            cells.extend(row["id"] for row in page["cells"])
            if not page["cell_page"]["has_more"]:
                break
            offset = int(page["cell_page"]["next_cursor"])
        assert (
            len(found_recipes) == len(set(found_recipes)) == 105 and set(found_recipes) == recipes
        )
        assert len(found_runs) == len(set(found_runs)) == 105 and set(found_runs) == runs
        assert len(cells) == len(set(cells)) == 106
        assert self.records(dataset, source["id"])[0]["text"] == "original evidence"
        self.history_dataset = dataset
        return {"executed_recipes": 105, "terminal_runs": 105, "readable_cells": 106}

    def permissions_for_every_workshop_tool(self):
        dataset = self.small
        source = self.small_source
        recipe = self.call(
            "save_dataset_pipeline",
            {
                "dataset": dataset,
                "name": "Permission fixture",
                "request_key": "permission",
                "steps": [{"operation": "select", "columns": ["text"]}],
            },
        )["pipeline"]
        capability = Capability(
            id=uuid.uuid4(), project=self.project, name="Access fixture", slug="access-fixture"
        )
        Capability.objects.bulk_create([capability])
        pipeline_run = DatasetPipelineRun.objects.filter(dataset_id=dataset).first()
        partition = DataPartitionPlan.objects.filter(project=self.project).first()
        payloads = {
            "start_dataset": {"brief": "Not authorized"},
            "cancel_dataset": {"dataset": dataset},
            "list_datasets": {},
            "inspect_dataset": {"dataset": dataset},
            "query_dataset": {"dataset": dataset, "sql": "SELECT * FROM t"},
            "create_dataset_from_traces": {"name": "Denied", "trace_ids": ["a" * 32]},
            "create_dataset_from_llm_calls": {
                "name": "Denied",
                "capability": str(capability.pk),
                "since": "2020-01-01T00:00:00Z",
            },
            "inspect_dataset_workbench": {"dataset": dataset},
            "save_dataset_pipeline": {
                "dataset": dataset,
                "name": "Denied",
                "request_key": "denied",
                "steps": [{"operation": "select", "columns": ["text"]}],
            },
            "run_dataset_pipeline": {
                "dataset": dataset,
                "pipeline": recipe["id"],
                "source_cell": source["id"],
                "source_fingerprint": source["fingerprint"],
                "request_key": "denied",
            },
            "import_dataset_version": {
                "dataset": dataset,
                "source_cell": source["id"],
                "source_fingerprint": source["fingerprint"],
                "request_key": "denied",
                "name": "Denied",
                "provenance": "Denied fixture",
                "imported_rows": [{"source_row": 0, "text": "Denied"}],
            },
            "update_dataset": {"dataset": dataset, "name": "Denied"},
            "explore_dataset": {
                "source_cell": source["id"],
                "name": "Denied",
                "request_key": "denied",
            },
            "derive_dataset": {
                "source_cell": source["id"],
                "name": "Denied",
                "request_key": "denied",
            },
            "create_data_partition": {
                "source_cell": source["id"],
                "name": "Denied",
                "request_key": "denied",
                "recipe": {"seed": 1, "fractions": {"train": 0.8, "final": 0.2}},
            },
            "retry_data_partition": {"partition": str(partition.pk)},
            "get_job": {"kind": "dataset_pipeline", "id": str(pipeline_run.pk)},
            "list_model_workflows": {"kind": "data_partition"},
            "list_projects": {},
        }
        readonly = {tool["name"] for tool in self.rpc("tools/list", key=self.read_key)["tools"]}
        counts = Counter()
        for name, args in payloads.items():
            if name not in readonly:
                result = self.call(name, args, key=self.read_key, error=True)
                assert result["error"]["code"] == "permission_denied", (name, result)
                counts["read_only_denials"] += 1
            else:
                self.call(name, args, key=self.read_key)
            if name != "list_projects":
                result = self.call(
                    name,
                    {**args, "project_id": str(self.foreign.pk)},
                    key=self.account_key,
                    error=True,
                )
                assert result["error"]["code"] == "project_required", (name, result)
                counts["foreign_project_denials"] += 1
        assert set(payloads) == EXPECTED
        foreign = Dataset.objects.create(
            project=self.foreign, name="Foreign fixture", state=Dataset.State.IDLE
        )
        for name in (
            "inspect_dataset",
            "query_dataset",
            "inspect_dataset_workbench",
            "save_dataset_pipeline",
            "run_dataset_pipeline",
            "import_dataset_version",
            "update_dataset",
            "cancel_dataset",
        ):
            self.call(name, {**payloads[name], "dataset": str(foreign.pk)}, error=True)
            counts["foreign_dataset_denials"] += 1
        for uri in (
            f"overmind://datasets/{foreign.pk}",
            f"overmind://jobs/dataset_run/{foreign.pk}",
        ):
            response = self.http.post(
                BASE + "/api/mcp/",
                headers={"X-Api-Key": self.key, "Accept": "application/json"},
                json={
                    "jsonrpc": "2.0",
                    "id": 1,
                    "method": "resources/read",
                    "params": {"uri": uri},
                },
            )
            assert response.status_code == 200 and "error" in response.json(), response.text[:300]
        response = self.http.get(
            BASE + f"/api/datasets/{foreign.pk}/export/", headers={"X-Api-Key": self.key}
        )
        assert response.status_code in {403, 404}
        assert self.inspect(dataset)["name"] != "Denied"
        return dict(counts)

    def cancel_actual_worker(self):
        dataset = self.large_dataset
        source = self.active(dataset)
        if source["rows"] < 50000:
            return {
                "not_exercised": "Run with at least 100,000 source rows to observe live cancellation"
            }
        recipe = self.call(
            "save_dataset_pipeline",
            {
                "dataset": dataset,
                "name": "Cancel live work",
                "request_key": "cancel-worker",
                "steps": [{"operation": "select", "columns": ["messages", "value"]}],
            },
        )["pipeline"]
        args = {
            "dataset": dataset,
            "pipeline": recipe["id"],
            "source_cell": source["id"],
            "source_fingerprint": source["fingerprint"],
            "request_key": "cancel-worker",
        }
        receipt = self.call("run_dataset_pipeline", args)
        identifier = receipt["run"]["id"]
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            job = self.call("get_job", {"kind": "dataset_pipeline", "id": identifier})
            if job["status"] == "running":
                break
            assert job["status"] == "queued", job
            time.sleep(0.02)
        assert job["status"] == "running", job
        self.call("run_dataset_pipeline", {**args, "request_key": "busy-request"}, error=True)
        self.call("update_dataset", {"dataset": dataset, "active": source["id"]}, error=True)
        self.call("cancel_dataset", {"dataset": dataset})
        self.poll("dataset_pipeline", identifier, expected="cancelled")
        assert self.active(dataset) == source
        assert self.call("run_dataset_pipeline", args)["run"]["state"] == "cancelled"
        replacement = self.call("run_dataset_pipeline", {**args, "request_key": "after-cancel"})
        self.poll("dataset_pipeline", replacement["run"]["id"])
        assert self.active(dataset)["rows"] == source["rows"]
        assert (
            self.resource(f"overmind://jobs/dataset_pipeline/{identifier}")["output_cell"] is None
        )
        return {
            "cancelled_state_observed": "running",
            "rows_in_flight": source["rows"],
            "replacement_completed": True,
        }

    def file_transport_formats(self):
        fixtures = {
            "table.csv": b'id,text\n007,"caf\xc3\xa9\nsecond line"\n008,"quoted ""answer"""\n',
            "table.tsv": b"id\ttext\n007\tcaf\xc3\xa9\n008\tsecond\n",
            "table.json": json.dumps(
                {"rows": [{"id": "007", "nested": [1, None, True]}, {"id": "008", "nested": [2]}]}
            ).encode(),
            "table.ndjson": b'{"id":"007","nested":{"x":1}}\n{"id":"008","nested":{"x":2}}\n',
            "table.jsonl.gz": gzip.compress(
                b'{"id":"007","text":"first"}\n{"id":"008","text":"second"}\n'
            ),
        }
        for filename, content in fixtures.items():
            (self.directory / filename).write_bytes(content)
        parquet = self.directory / "table.parquet"
        pq.write_table(
            pa.Table.from_pylist(
                [{"id": "007", "nested": {"x": 1}}, {"id": "008", "nested": {"x": 2}}]
            ),
            parquet,
        )
        for filename in [*fixtures, "table.parquet"]:
            result = upload_file(
                self.directory / filename,
                project_id=str(self.project.pk),
                api_key=self.key,
                api_url=BASE,
                intent="explore",
            )
            dataset = result["id"]
            wait_until_ready(dataset, api_key=self.key, api_url=BASE, poll=0.25, timeout=600)
            source = self.active(dataset)
            assert source["rows"] == 2, (filename, source)
            actual = self.records(dataset)
            assert [row["id"] for row in actual] == ["007", "008"], (filename, actual)
            self.pipeline(dataset, source, [{"operation": "select", "columns": ["id"]}])
            exported = self.directory / f"export-{filename}.jsonl"
            export_dataset(
                dataset,
                cell=self.active(dataset)["id"],
                output=exported,
                api_key=self.key,
                api_url=BASE,
            )
            assert [json.loads(line)["id"] for line in exported.read_text().splitlines()] == [
                "007",
                "008",
            ]
        malformed = {
            "duplicate.csv": b"name,Name\nfirst,second\n",
            "broken.jsonl": b'{"text":"valid"}\n{"broken":\n',
            "empty.jsonl": b"",
        }
        rejected = 0
        for filename, content in malformed.items():
            path = self.directory / filename
            path.write_bytes(content)
            try:
                result = upload_file(
                    path, project_id=str(self.project.pk), api_key=self.key, api_url=BASE
                )
            except Exception as exc:
                assert "500" not in str(exc), str(exc)
                rejected += 1
            else:
                self.poll("dataset_run", result["id"], expected="error")
                assert self.active(result["id"]) is None
                rejected += 1
        return {"file_formats_round_tripped": 6, "invalid_files_rejected": rejected}

    def document_source_evidence(self):
        fixture_directory = Path(os.environ["WORKSHOP_ACCEPTANCE_DOCUMENTS"])
        (self.directory / "handbook.txt").write_text(
            "Service handbook\n\nReturns are accepted for thirty days.\n"
        )
        (self.directory / "handbook.md").write_text(
            "# Service handbook\n\nReturns are accepted for thirty days.\n"
        )
        document = Document()
        document.add_paragraph("Returns are accepted for thirty days.")
        document.save(self.directory / "handbook.docx")
        image_paths = []
        with pdfium.PdfDocument(fixture_directory / "scanned.pdf") as pdf:
            page = pdf[0]
            bitmap = page.render(scale=2)
            raster = bitmap.to_pil().convert("RGB")
            for extension, format_name in (("png", "PNG"), ("jpg", "JPEG"), ("webp", "WEBP")):
                path = self.directory / f"scan.{extension}"
                raster.save(path, format=format_name)
                image_paths.append(path)
            bitmap.close()
            page.close()
        files = [self.directory / name for name in ("handbook.txt", "handbook.md", "handbook.docx")]
        files += [fixture_directory / name for name in ("scanned.pdf", "mixed.pdf")]
        files += image_paths
        for path in files:
            result = upload_file(
                path,
                project_id=str(self.project.pk),
                api_key=self.key,
                api_url=BASE,
                intent="explore",
            )
            dataset = result["id"]
            wait_until_ready(dataset, api_key=self.key, api_url=BASE, poll=0.25, timeout=600)
            records = self.records(dataset)
            assert "Returns are accepted for thirty days." in " ".join(
                row["text"] for row in records
            ), path.name
            sha = hashlib.sha256(path.read_bytes()).hexdigest()
            for row in records:
                assert row["_overmind_document_id"] == sha
                assert row["_overmind_provenance"]["evidence"][0]["document_id"] == sha
                assert "messages" not in row
            detail = self.inspect(dataset)
            artifact = detail["sources"][0]
            downloaded = self.http.get(
                BASE + f"/api/datasets/{dataset}/sources/{artifact['id']}/",
                headers={"X-Api-Key": self.key},
            )
            assert downloaded.status_code == 200 and downloaded.content == path.read_bytes(), (
                path.name,
                artifact,
                downloaded.status_code,
                downloaded.text[:300] if downloaded.status_code != 200 else "Bytes differ",
            )
            source = self.active(dataset)
            self.pipeline(dataset, source, [{"operation": "select", "columns": ["text"]}])
            assert all(row["_overmind_provenance"]["evidence"] for row in self.records(dataset))
        return {
            "document_and_image_uploads": len(files),
            "original_bytes_verified": len(files),
            "evidence_after_projection": len(files),
        }

    def pdf_poll(self, dataset, expected="idle"):
        started = time.monotonic()
        observations = []
        previous = None
        polls = 0
        while time.monotonic() - started < 1500:
            job = self.call("get_job", {"kind": "dataset_run", "id": dataset})
            polls += 1
            internal = Dataset.objects.get(id=dataset, project=self.project)
            stored_progress = internal.source_spec.get("landing_progress")
            if stored_progress and job["status"] == "landing":
                exposed = job["progress"].get("landing")
                if exposed is not None:
                    assert exposed["total"] == stored_progress["total"]
                    assert 0 <= exposed["completed"] <= exposed["total"]
            state = {
                "status": job["status"],
                "mcp_progress": job.get("progress"),
                "stored_landing_progress": internal.source_spec.get("landing_progress"),
            }
            if state != previous:
                observations.append({"seconds": round(time.monotonic() - started, 3), **state})
                previous = state
            if job["status"] not in {"landing", "running", "queued"}:
                assert job["status"] == expected, job
                return {
                    "landing_seconds": round(time.monotonic() - started, 3),
                    "polls": polls,
                    "observations": observations,
                    "error": job.get("job_error"),
                }
            time.sleep(0.5)
        raise AssertionError(f"PDF landing exceeded 1500 seconds: {dataset}")

    def pdf_stage(self, path):
        headers = {"X-Api-Key": self.key}
        response = self.http.post(
            BASE + "/api/uploads/", headers=headers, json={"filename": path.name}
        )
        assert response.status_code == 201, response.text[:500]
        reserved = response.json()
        identifier = reserved["upload_id"]
        self.pdf_uploads.append(identifier)
        offset = 0
        with path.open("rb") as stream:
            while chunk := stream.read(reserved["chunk_bytes"]):
                response = self.http.put(
                    BASE + f"/api/uploads/{identifier}/chunk/",
                    headers={**headers, "Content-Type": "application/octet-stream"},
                    params={"offset": offset},
                    content=chunk,
                )
                assert response.status_code == 200, response.text[:500]
                offset += len(chunk)
                assert response.json()["received"] == offset
        return identifier

    def pdf_batch_submit(self, uploads, dataset=None, expected_status=201):
        body = {"uploads": uploads}
        endpoint = f"/api/datasets/{dataset}/source/" if dataset else "/api/datasets/"
        if not dataset:
            body = {
                "project": str(self.project.pk),
                "name": "PDF batch regression",
                "intent": "explore",
                "capability": None,
                "source": body,
            }
        response = self.http.post(BASE + endpoint, headers={"X-Api-Key": self.key}, json=body)
        assert response.status_code == expected_status, (response.status_code, response.text[:500])
        return response.json()

    def pdf_case(self, fixture):
        directory = Path(os.environ["WORKSHOP_ACCEPTANCE_PDFS"])
        path = directory / fixture["filename"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == fixture["sha256"]
        started = time.monotonic()
        if path.stat().st_size > 100 * 1024**2:
            before = Dataset.objects.filter(project=self.project).count()
            reserved = self.http.post(
                BASE + "/api/uploads/",
                headers={"X-Api-Key": self.key},
                json={"filename": path.name},
            )
            assert reserved.status_code == 201
            identifier = reserved.json()["upload_id"]
            self.pdf_uploads.append(identifier)
            assert reserved.json()["max_bytes"] == 100 * 1024**2, reserved.json()
            rejected = self.http.put(
                BASE + f"/api/uploads/{identifier}/chunk/",
                headers={"X-Api-Key": self.key, "Content-Type": "application/octet-stream"},
                params={"offset": 100 * 1024**2},
                content=b"x",
            )
            assert rejected.status_code == 409 and "104857600" in rejected.text, (
                rejected.status_code,
                rejected.text,
            )
            assert files.upload_received(identifier) == 0
            assert Dataset.objects.filter(project=self.project).count() == before
            requests_sent = []

            def capture(response, *args, **kwargs):
                requests_sent.append(response.request.method)
                if response.request.method == "POST" and response.status_code == 201:
                    self.pdf_uploads.append(response.json()["upload_id"])

            with requests.Session() as session:
                session.hooks["response"].append(capture)
                try:
                    upload_file(
                        path,
                        project_id=str(self.project.pk),
                        api_key=self.key,
                        api_url=BASE,
                        intent="explore",
                        session=session,
                    )
                except DatasetUploadError as exc:
                    assert "104857600" in str(exc), str(exc)
                else:
                    raise AssertionError("SDK accepted a PDF over the advertised byte limit")
            assert requests_sent == ["POST"], requests_sent
            assert Dataset.objects.filter(project=self.project).count() == before
            return {**fixture, "rejected_before_transfer": True, "uploaded_bytes": 0}
        result = upload_file(
            path, project_id=str(self.project.pk), api_key=self.key, api_url=BASE, intent="explore"
        )
        dataset = result["id"]
        timing = self.pdf_poll(dataset, fixture["expected"])
        detail = self.inspect(dataset)
        facts = {
            **fixture,
            **timing,
            "source_to_ready_seconds": round(time.monotonic() - started, 3),
        }
        if fixture["expected"] == "error":
            assert detail["active"] is None and timing["error"]
            if fixture["pages"] and fixture["pages"] > 2000:
                assert "2000" in timing["error"] and "2001" in timing["error"], timing
            assert not Dataset.objects.get(id=dataset).cells.filter(state="ok").exists()
            upload_file(
                directory / "native-1.pdf",
                project_id=str(self.project.pk),
                api_key=self.key,
                api_url=BASE,
                dataset=dataset,
            )
            self.pdf_poll(dataset)
            assert self.active(dataset)["rows"] > 0
            return {**facts, "recovered_with_valid_source": True}
        records = self.records(dataset)
        by_page = {}
        for row in records:
            assert row["_overmind_document_id"] == fixture["sha256"]
            evidence = row["_overmind_provenance"]["evidence"]
            assert evidence and all(e["document_id"] == fixture["sha256"] for e in evidence)
            assert all(e["page"] == row["page"] and e["regions"] for e in evidence)
            assert "messages" not in row
            by_page.setdefault(row["page"], []).append(row["text"])
        assert set(by_page) == set(range(1, fixture["pages"] + 1)), sorted(by_page)
        for number, texts in by_page.items():
            text = " ".join(texts)
            if fixture["kind"] == "native":
                assert f"{fixture['marker']} page {number:04d}" in text, (number, text)
            elif fixture["kind"] == "scanned" or number % 3 != 1:
                assert "Returns are accepted for thirty days." in text, (number, text)
        artifact = detail["sources"][0]
        extraction = artifact["extraction"]
        assert extraction["pages"] == fixture["pages"]
        assert extraction["method"] and extraction["version"] and extraction["limitations"]
        if fixture["kind"] in {"scanned", "mixed"}:
            assert extraction["ocr"]["engine"] == "tesseract"
            assert extraction["ocr"]["languages"] == ["eng"]
            assert extraction["ocr"]["version"] and extraction["ocr"]["pages"]
        if fixture["filename"] == "scanned-25.pdf":
            ocr_observations = [
                item["mcp_progress"]["landing"]
                for item in timing["observations"]
                if (item["mcp_progress"].get("landing") or {}).get("stage") == "ocr"
            ]
            assert ocr_observations and all(item["pages_total"] == 25 for item in ocr_observations)
            assert max(item["ocr_pages_completed"] for item in ocr_observations) > 0
        response = self.http.get(
            BASE + f"/api/datasets/{dataset}/sources/{artifact['id']}/",
            headers={"X-Api-Key": self.key},
        )
        assert response.status_code == 200
        assert hashlib.sha256(response.content).hexdigest() == fixture["sha256"]
        source = detail["active"]
        _, _, job = self.pipeline(dataset, source, [{"operation": "select", "columns": ["text"]}])
        output = self.records(dataset)
        assert [row["text"] for row in output] == [row["text"] for row in records]
        assert [row["_overmind_provenance"]["evidence"] for row in output] == [
            row["_overmind_provenance"]["evidence"] for row in records
        ]
        assert self.records(dataset, source["id"]) == records
        receipt = self.resource(job["resource"]["uri"])
        assert receipt["result"]["semantic_quality"] == "unmeasured"
        return {
            **facts,
            "rows": len(records),
            "pages_with_evidence": len(by_page),
            "extraction": extraction,
            "mcp_extraction": artifact["extraction"],
            "catalogue_gap": None,
            "semantic_quality": receipt["result"]["semantic_quality"],
            "original_bytes_and_prior_cell_verified": True,
        }

    def pdf_quantity(self, count):
        directory = Path(os.environ["WORKSHOP_ACCEPTANCE_PDFS"])
        uploads = [
            self.pdf_stage(directory / f"batch-{index % 25:02d}.pdf") for index in range(count)
        ]
        dataset = self.pdf_batch_submit(uploads)["id"]
        timing = self.pdf_poll(dataset)
        detail = self.inspect(dataset)
        records = self.records(dataset)
        row_text = Counter(row["text"] for row in records)
        for index in range(min(count, 25)):
            assert (
                row_text[f"Batch document {index:02d} page 0001"] == 1 + (count - 1 - index) // 25
            )
        assert detail["sources_total"] == count
        sources = self.source_inventory(dataset)
        assert len(sources) == count and len({source["id"] for source in sources}) == min(count, 25)
        for source in sources:
            response = self.http.get(
                BASE + f"/api/datasets/{dataset}/sources/{source['id']}/",
                headers={"X-Api-Key": self.key},
            )
            assert response.status_code == 200
            assert hashlib.sha256(response.content).hexdigest() == source["sha256"]
        return {
            **timing,
            "uploaded_files": count,
            "unique_document_contents": min(count, 25),
            "original_downloads_verified": count,
            "rows": len(records),
            "mcp_sources_visible": len(sources),
            "mcp_sources_total": detail["sources_total"],
            "catalogue_gap": None,
        }

    def pdf_mixed_batch(self):
        directory = Path(os.environ["WORKSHOP_ACCEPTANCE_PDFS"])
        names = {
            "native-1.pdf",
            "native-10.pdf",
            "native-100.pdf",
            "native-500.pdf",
            "native-2000.pdf",
            "scanned-1.pdf",
            "scanned-25.pdf",
            "mixed-30.pdf",
            "large-scanned.pdf",
            "batch-00.pdf",
        }
        fixtures = [
            fixture
            for fixture in json.loads((directory / "manifest.json").read_text())
            if fixture["filename"] in names
        ]
        uploads = [self.pdf_stage(directory / fixture["filename"]) for fixture in fixtures]
        dataset = self.pdf_batch_submit(uploads)["id"]
        timing = self.pdf_poll(dataset)
        records = self.records(dataset)
        pages_by_document = {}
        for row in records:
            identity = row["_overmind_document_id"]
            pages_by_document.setdefault(identity, set()).add(row["page"])
            assert all(
                evidence["document_id"] == identity and evidence["page"] == row["page"]
                for evidence in row["_overmind_provenance"]["evidence"]
            )
        assert set(pages_by_document) == {fixture["sha256"] for fixture in fixtures}
        for fixture in fixtures:
            assert pages_by_document[fixture["sha256"]] == set(range(1, fixture["pages"] + 1))
        source = self.active(dataset)
        self.pipeline(dataset, source, [{"operation": "select", "columns": ["text"]}])
        output = self.records(dataset)
        assert [row["text"] for row in output] == [row["text"] for row in records]
        assert [row["_overmind_provenance"]["evidence"] for row in output] == [
            row["_overmind_provenance"]["evidence"] for row in records
        ]
        assert self.records(dataset, source["id"]) == records
        return {
            **timing,
            "files": len(fixtures),
            "pages": sum(fixture["pages"] for fixture in fixtures),
            "bytes": sum(fixture["bytes"] for fixture in fixtures),
            "rows": len(records),
            "per_document_page_identity_and_transformed_evidence_verified": True,
        }

    def pdf_atomic_batch_recovery(self):
        directory = Path(os.environ["WORKSHOP_ACCEPTANCE_PDFS"])
        dataset = upload_file(
            directory / "native-1.pdf",
            project_id=str(self.project.pk),
            api_key=self.key,
            api_url=BASE,
            intent="explore",
        )["id"]
        self.pdf_poll(dataset)
        source = self.active(dataset)
        original = self.records(dataset)
        uploads = [
            self.pdf_stage(directory / name)
            for name in ("batch-00.pdf", "corrupt.pdf", "batch-01.pdf")
        ]
        self.pdf_batch_submit(uploads, dataset, expected_status=202)
        failed = self.pdf_poll(dataset, "error")
        assert self.active(dataset) == source
        assert self.records(dataset) == original
        uploads = [self.pdf_stage(directory / name) for name in ("batch-00.pdf", "batch-01.pdf")]
        self.pdf_batch_submit(uploads, dataset, expected_status=202)
        self.pdf_poll(dataset)
        assert self.records(dataset, source["id"]) == original
        assert self.active(dataset)["rows"] > source["rows"]
        assert self.inspect(dataset)["sources_total"] == 3
        return {"failed_attempt": failed, "previous_cell_preserved": True, "recovered_sources": 3}

    def pdf_batch_limit(self):
        directory = Path(os.environ["WORKSHOP_ACCEPTANCE_PDFS"])
        before = Dataset.objects.filter(project=self.project).count()
        uploads = [self.pdf_stage(directory / "native-1.pdf") for _ in range(101)]
        result = self.pdf_batch_submit(uploads, expected_status=400)
        assert "100" in json.dumps(result), result
        assert Dataset.objects.filter(project=self.project).count() == before
        return {"files": 101, "rejected_without_dataset": True, "error": result}

    def source_bindings_and_metadata(self):
        dataset = self.upload("bindings.jsonl", [{"text": "first"}, {"text": "second"}])
        source = self.active(dataset)
        artifact = self.upload(
            "artifact.jsonl", [{"text": "Derived first", "_overmind_parent_rows": [0]}]
        )
        artifact_cell = self.active(artifact)
        foreign = Dataset.objects.create(
            project=self.foreign, name="Foreign cells", state=Dataset.State.IDLE
        )
        foreign_cell = Cell.objects.create(
            dataset=foreign, position=0, state="ok", fingerprint="f" * 64, rows=1
        )
        arguments = {
            "dataset": dataset,
            "source_cell": source["id"],
            "source_fingerprint": source["fingerprint"],
            "artifact_cell": artifact_cell["id"],
            "artifact_fingerprint": artifact_cell["fingerprint"],
            "name": "Pinned artifact",
            "request_key": "artifact-bindings",
            "provenance": "Native fixture transformation",
        }
        for changes in (
            {"artifact_fingerprint": "0" * 64},
            {
                "artifact_cell": str(foreign_cell.pk),
                "artifact_fingerprint": foreign_cell.fingerprint,
            },
            {
                "source_cell": artifact_cell["id"],
                "source_fingerprint": artifact_cell["fingerprint"],
            },
            {"imported_rows": [{"source_row": 0, "text": "Ambiguous input"}]},
        ):
            self.call("import_dataset_version", {**arguments, **changes}, error=True)
            assert self.active(dataset) == source
        for tool in ("explore_dataset", "derive_dataset", "create_data_partition"):
            args = {
                "source_cell": str(foreign_cell.pk),
                "name": "Foreign source",
                "request_key": "foreign-source",
            }
            if tool == "create_data_partition":
                args["recipe"] = {"seed": 1, "fractions": {"train": 0.8, "final": 0.2}}
            self.call(tool, args, error=True)
        receipt = self.call("import_dataset_version", arguments)
        self.poll("dataset_pipeline", receipt["run"]["id"])
        imported = self.active(dataset)
        response = self.http.delete(
            BASE + f"/api/datasets/{artifact}/", headers={"X-Api-Key": self.key}
        )
        assert response.status_code == 409, (response.status_code, response.text[:500])
        assert self.active(artifact) == artifact_cell
        self.call(
            "update_dataset",
            {
                "dataset": dataset,
                "active": source["id"],
                "name": "Selected source",
                "intent": "train",
            },
        )
        assert self.active(dataset)["id"] == source["id"]
        assert self.active(dataset)["fingerprint"] == source["fingerprint"]
        assert self.inspect(dataset)["name"] == "Selected source"
        assert self.inspect(dataset)["intent"] == "train"
        self.call("update_dataset", {"dataset": dataset, "active": artifact_cell["id"]}, error=True)
        self.call("update_dataset", {"dataset": dataset, "active": None})
        assert self.active(dataset)["id"] == imported["id"]
        assert self.active(dataset)["fingerprint"] == imported["fingerprint"]
        operation = self.call(
            "explore_dataset",
            {"source_cell": imported["id"], "name": "Pin for exploration", "request_key": "pin"},
        )
        self.poll("data_exploration", operation["workflow"]["id"])
        self.call(
            "update_dataset",
            {"dataset": dataset, "name": "Must roll back", "intent": "eval"},
            error=True,
        )
        assert self.inspect(dataset)["name"] == "Selected source"
        assert self.inspect(dataset)["intent"] == "train"
        self.call("update_dataset", {"dataset": dataset, "capability": None}, error=True)
        self.call("update_dataset", {"dataset": dataset, "active": source["id"]})
        assert self.records(dataset, imported["id"])[0]["text"] == "Derived first"
        for name in ("message_dataset_agent", "run_dataset", "manage_dataset_workflow"):
            result = self.call(name, {"dataset": dataset}, error=True)
            assert result["error"]["code"] == "invalid_tool"
        return {
            "stale_foreign_or_ambiguous_imports_rejected": 4,
            "foreign_source_workflows_rejected": 3,
            "artifact_retained": True,
            "frozen_metadata_atomic": True,
            "removed_tools_rejected": 3,
        }

    def partition_name_boundary(self):
        dataset = self.upload("long-partition.jsonl", [{"text": f"Row {i}"} for i in range(40)])
        source = self.active(dataset)
        names = ["n" * 255, "界" * 255]
        for index, name in enumerate(names):
            args = {
                "source_cell": source["id"],
                "name": name,
                "request_key": f"long-partition-{index}",
                "recipe": {"seed": 7, "fractions": {"train": 0.5, "development": 0.5}},
            }
            receipt = self.call("create_data_partition", args)
            identifier = receipt["workflow"]["id"]
            deadline = time.monotonic() + 60
            while time.monotonic() < deadline:
                detail = self.resource(f"overmind://jobs/data_partition/{identifier}")
                if detail["state"] in {"failed", "completed"}:
                    assert detail["state"] == "completed", detail
                    break
                time.sleep(0.25)
            else:
                raise AssertionError("Maximum-length partition timed out")
            self.poll("data_partition", identifier)
            detail = self.resource(f"overmind://jobs/data_partition/{identifier}")
            assert detail["name"] == name
            assert len(detail["members"]) == 2
            for member in detail["members"]:
                member_name = self.inspect(member["cell"]["dataset_id"])["name"]
                assert len(member_name) <= 255 and member_name.endswith(f" · {member['role']}")
            assert self.call("create_data_partition", args)["workflow"]["id"] == identifier
        return {"maximum_length_names": len(names), "roles_per_name": 2}

    def partition_declared_groups(self, field="cohort"):
        dataset = self.upload(
            "explicit-groups.jsonl",
            [{"text": f"Distinct {i}", field: f"group-{i // 4}"} for i in range(40)],
        )
        source = self.active(dataset)
        args = {
            "source_cell": source["id"],
            "name": "Declared groups",
            "request_key": f"misspelled-group-{field}",
            "recipe": {
                "seed": 11,
                "fractions": {"train": 0.5, "final": 0.5},
                "group_by": ["cohrot"],
            },
        }
        bad = self.call("create_data_partition", args)["workflow"]["id"]
        self.poll("data_partition", bad, expected="failed")
        detail = self.resource(f"overmind://jobs/data_partition/{bad}")
        assert "cohrot" in detail["error"] and not detail["members"]
        assert self.active(dataset) == source
        good = self.call(
            "create_data_partition",
            {
                **args,
                "request_key": f"declared-group-{field}",
                "recipe": {**args["recipe"], "group_by": [field]},
            },
        )["workflow"]["id"]
        self.poll("data_partition", good)
        detail = self.resource(f"overmind://jobs/data_partition/{good}")
        member = detail["members"][0]["cell"]
        child = member["dataset_id"]
        self.pipeline(child, self.active(child), [{"operation": "select", "columns": ["text"]}])
        projected = self.call(
            "create_data_partition",
            {
                **args,
                "source_cell": self.active(child)["id"],
                "request_key": f"projected-group-{field}",
                "recipe": {**args["recipe"], "group_by": [field]},
            },
        )["workflow"]["id"]
        self.poll("data_partition", projected)
        owner = {}
        for member in self.resource(f"overmind://jobs/data_partition/{projected}")["members"]:
            for row in self.records(member["cell"]["dataset_id"]):
                group = int(row["text"].removeprefix("Distinct ")) // 4
                assert owner.setdefault(group, member["role"]) == member["role"]
        return {
            "field": field,
            "missing_group_rejected": True,
            "projected_lineage_groups": len(owner),
        }

    def account_mutations_and_resources(self):
        ProjectMembership.objects.get_or_create(user=self.user, project=self.foreign)
        datasets = []
        followed = 0
        for project in (self.project, self.foreign):

            def call(name, args=None, project=project):
                return self.call(
                    name, {**(args or {}), "project_id": str(project.pk)}, key=self.account_key
                )

            draft = call("start_dataset", {"brief": "Account-key regression", "intent": "explore"})
            dataset = draft["dataset"]["id"]
            path = self.directory / f"account-{project.pk}.jsonl"
            path.write_text(json.dumps({"text": str(project.pk), "value": 1}) + "\n")
            upload_file(
                path,
                project_id=str(project.pk),
                api_key=self.account_key,
                api_url=BASE,
                dataset=dataset,
            )
            wait_until_ready(dataset, api_key=self.account_key, api_url=BASE, poll=0.25, timeout=60)
            source = call("inspect_dataset", {"dataset": dataset})["active"]
            recipe = call(
                "save_dataset_pipeline",
                {
                    "dataset": dataset,
                    "name": "Account selection",
                    "request_key": "same-key",
                    "steps": [{"operation": "select", "columns": ["text"]}],
                },
            )
            receipt = call(
                "run_dataset_pipeline",
                {
                    "dataset": dataset,
                    "pipeline": recipe["pipeline"]["id"],
                    "source_cell": source["id"],
                    "source_fingerprint": source["fingerprint"],
                    "request_key": "same-key",
                },
            )
            deadline = time.monotonic() + 60
            while time.monotonic() < deadline:
                job = call("get_job", {"kind": "dataset_pipeline", "id": receipt["run"]["id"]})
                if job["status"] in {"completed", "failed", "cancelled"}:
                    assert job["status"] == "completed", job
                    break
                time.sleep(0.25)
            else:
                raise AssertionError("Account-key job timed out")
            for response in (draft, recipe, receipt, job):
                for link in response.get("resource_links", []):
                    assert f"project_id={project.pk}" in link["uri"], link
                    assert self.resource(link["uri"], key=self.account_key)
                    followed += 1
            call("update_dataset", {"dataset": dataset, "name": "Account-owned result"})
            result = call("query_dataset", {"dataset": dataset, "sql": "SELECT text FROM t"})
            assert result["rows"] == [{"text": str(project.pk)}]
            datasets.append(dataset)
        for project, other_dataset in ((self.project, datasets[1]), (self.foreign, datasets[0])):
            self.call(
                "inspect_dataset",
                {"dataset": other_dataset, "project_id": str(project.pk)},
                key=self.account_key,
                error=True,
            )
        return {
            "authorized_projects": 2,
            "successful_mutation_workflows": 2,
            "resource_links_followed": followed,
        }

    def partition_group_aliases(self):
        dataset = self.upload(
            "group-aliases.jsonl",
            [{"text": f"Alias {i}", "source_trace_id": f"trace-{i // 4}"} for i in range(40)],
        )
        source = self.active(dataset)
        merged = self.call(
            "import_dataset_version",
            {
                "dataset": dataset,
                "source_cell": source["id"],
                "source_fingerprint": source["fingerprint"],
                "request_key": "merge-trace-groups",
                "name": "Merged trace evidence",
                "provenance": "Eight consecutive source observations per result; all contributing parents retained.",
                "imported_rows": [
                    {"text": f"Merged {i}", "_overmind_parent_rows": list(range(i, i + 8))}
                    for i in range(0, 40, 8)
                ],
            },
        )
        self.poll("dataset_pipeline", merged["run"]["id"])
        assert all(not row.get("source_trace_id") for row in self.records(dataset))
        projected = self.active(dataset)
        args = {
            "source_cell": projected["id"],
            "name": "Projected trace groups",
            "request_key": "group-aliases",
            "recipe": {
                "seed": 5,
                "fractions": {"train": 0.5, "final": 0.5},
                "group_by": ["source_trace_id"],
            },
        }
        receipt = self.call("create_data_partition", args)
        self.poll("data_partition", receipt["workflow"]["id"])
        owner = {}
        for member in self.resource(f"overmind://jobs/data_partition/{receipt['workflow']['id']}")[
            "members"
        ]:
            for row in self.records(member["cell"]["dataset_id"]):
                groups = [
                    value
                    for key, value in row["_overmind_provenance"]["source_group_keys"]
                    if key == "trace_id"
                ]
                assert groups
                for group in groups:
                    assert owner.setdefault(group, member["role"]) == member["role"]
        assert len(owner) == 10
        bad = self.call(
            "create_data_partition",
            {
                **args,
                "request_key": "nonexistent-content-field",
                "recipe": {**args["recipe"], "group_by": ["content"]},
            },
        )
        self.poll("data_partition", bad["workflow"]["id"], expected="failed")
        active = self.active(dataset)
        assert {key: active[key] for key in ("id", "fingerprint", "rows")} == {
            key: projected[key] for key in ("id", "fingerprint", "rows")
        }, {"before": projected, "after": active}
        return {
            "projected_trace_groups": len(owner),
            "missing_content_field_rejected": True,
            "version_label_before_freeze": projected["version"],
            "version_label_after_freeze": active["version"],
            "selected_identity_and_content_unchanged": True,
        }

    def concurrent_partition_keys(self):
        sources = [
            self.active(
                self.upload(
                    f"partition-race-{i}.jsonl",
                    [{"text": f"Source {i} row {j}"} for j in range(10000)],
                )
            )
            for i in range(2)
        ]
        conflicts = 0
        for attempt in range(3):
            barrier = Barrier(2)

            def submit(source, barrier=barrier, attempt=attempt):
                barrier.wait(timeout=20)
                return self.rpc(
                    "tools/call",
                    {
                        "name": "create_data_partition",
                        "arguments": {
                            "source_cell": source["id"],
                            "name": "Concurrent plan",
                            "request_key": f"cross-source-{attempt}",
                            "recipe": {"seed": 12, "fractions": {"train": 0.8, "final": 0.2}},
                        },
                    },
                )

            with ThreadPoolExecutor(max_workers=2) as pool:
                responses = list(pool.map(submit, sources))
            successful = [
                value["structuredContent"] for value in responses if not value.get("isError")
            ]
            rejected = [value["structuredContent"] for value in responses if value.get("isError")]
            for value in successful:
                self.poll("data_partition", value["workflow"]["id"])
            assert len(successful) == len(rejected) == 1, responses
            assert rejected[0]["error"]["code"] != "internal_error", rejected
            assert (
                DataPartitionPlan.objects.filter(
                    project=self.project, request_key=f"cross-source-{attempt}"
                ).count()
                == 1
            )
            conflicts += 1
        return {"concurrent_pairs": 3, "stable_conflicts": conflicts, "source_rows_each": 10000}

    def catalogue_capability_audit(self):
        dataset = self.upload(
            "sample-source.jsonl", [{"text": f"Evidence {i}", "stratum": i % 5} for i in range(100)]
        )
        source = self.active(dataset)
        tools = {tool["name"]: tool for tool in self.rpc("tools/list")["tools"]}
        assert "sampling" not in tools["derive_dataset"]["inputSchema"]["properties"]
        self.call(
            "derive_dataset",
            {
                "source_cell": source["id"],
                "name": "Unsupported sample",
                "request_key": "unsupported-sample",
                "sampling": {"rows": 25, "seed": 7},
            },
            error=True,
        )
        original = self.records(dataset)
        rng = random.Random(7)
        selected = [
            row
            for stratum in range(5)
            for row in rng.sample([row for row in original if row["stratum"] == stratum], 5)
        ]
        imported = self.call(
            "import_dataset_version",
            {
                "dataset": dataset,
                "source_cell": source["id"],
                "source_fingerprint": source["fingerprint"],
                "request_key": "native-sample",
                "name": "Native stratified sample",
                "provenance": "Native Python random.Random(7), five observations per stratum, selected from the pinned 100-row source.",
                "imported_rows": selected,
            },
        )
        self.poll("dataset_pipeline", imported["run"]["id"])
        assert Counter(row["stratum"] for row in self.records(dataset)) == Counter(
            dict.fromkeys(range(5), 5)
        )
        inventory = None
        for i in range(12):
            inventory = self.upload(
                f"source-inventory-{i}.jsonl", [{"text": f"File {i}"}], dataset=inventory
            )
        detail = self.inspect(inventory)
        assert detail["sources_total"] == 12 and len(detail["sources"]) == 10
        assert "source_offset" in tools["inspect_dataset"]["inputSchema"]["properties"]
        all_sources = self.source_inventory(inventory)
        assert len(all_sources) == 12
        return {
            "native_sample_rows": 25,
            "sample_strata": 5,
            "source_files": 12,
            "mcp_visible_source_files": len(all_sources),
            "catalogue_gaps": [
                "Sampling allocation cannot be published by a deterministic platform recipe; native-client import works",
            ],
        }

    def large_partition_exact_membership(self):
        count = 100000
        dataset = self.upload(
            "large-partition.jsonl",
            (
                {"text": f"Unique evidence {i}", "group_id": f"cohort-{i // 5}", "value": i}
                for i in range(count)
            ),
        )
        source = self.active(dataset)
        receipt = self.call(
            "create_data_partition",
            {
                "source_cell": source["id"],
                "name": "Large grouped roles",
                "request_key": "large-grouped-roles",
                "recipe": {
                    "seed": 17,
                    "fractions": {
                        "train": 0.7,
                        "development": 0.1,
                        "calibration": 0.1,
                        "final": 0.1,
                    },
                    "group_by": ["group_id"],
                },
            },
        )
        self.poll("data_partition", receipt["workflow"]["id"])
        result = self.resource(f"overmind://jobs/data_partition/{receipt['workflow']['id']}")
        observed = set()
        owners = {}
        roles = {}
        for member in result["members"]:
            cell = member["cell"]
            destination = self.directory / f"large-partition-{member['role']}.jsonl"
            export_dataset(
                cell["dataset_id"],
                cell=cell["id"],
                output=destination,
                api_key=self.key,
                api_url=BASE,
            )
            rows = 0
            with destination.open() as records:
                for line in records:
                    row = json.loads(line)
                    value = row["value"]
                    assert value not in observed and row["text"] == f"Unique evidence {value}"
                    assert owners.setdefault(row["group_id"], member["role"]) == member["role"]
                    assert row["_overmind_provenance"]["partition"]["source_cell"] == source["id"]
                    observed.add(value)
                    rows += 1
            assert rows == cell["rows"] == result["report"]["counts"][member["role"]]
            roles[member["role"]] = rows
        assert observed == set(range(count)) and len(owners) == count // 5
        assert self.active(dataset) == source
        return {
            "source_rows": count,
            "exported_rows_verified": len(observed),
            "disjoint_groups": len(owners),
            "role_rows": roles,
        }

    def cleanup(self):
        projects = [self.project.pk, self.foreign.pk]
        APIToken.objects.filter(user=self.user).delete()
        active = Dataset.objects.filter(
            project_id__in=projects, state__in=["landing", "running"]
        ).exists() or any(
            model.objects.filter(project_id__in=projects, state__in=["queued", "running"]).exists()
            for model in (DataExploration, DataPartitionPlan)
        )
        if active:
            print(
                json.dumps(
                    {
                        "cleanup": "fixtures retained while workers are active",
                        "projects": [str(pk) for pk in projects],
                    }
                ),
                flush=True,
            )
            return
        for upload_id in self.pdf_uploads:
            files.discard_upload(upload_id)
        directories = [
            paths.dataset_dir(pk)
            for pk in Dataset.objects.filter(project_id__in=projects).values_list("pk", flat=True)
        ]
        directories += [
            paths.media_root() / "explorations" / str(pk)
            for pk in DataExploration.objects.filter(project_id__in=projects).values_list(
                "pk", flat=True
            )
        ]
        directories += [
            paths.media_root() / "data-partitions" / str(pk)
            for pk in DataPartitionPlan.objects.filter(project_id__in=projects).values_list(
                "pk", flat=True
            )
        ]
        DataExploration.objects.filter(project_id__in=projects).delete()
        DataPartitionPlan.objects.filter(project_id__in=projects).delete()
        DatasetPipelineRun.objects.filter(dataset__project_id__in=projects).delete()
        Project.objects.filter(pk__in=projects).delete()
        self.user.delete()
        for directory in directories:
            if directory.is_dir():
                shutil.rmtree(directory)
        self.http.close()
        print(
            json.dumps(
                {
                    "cleanup": "removed only this replay's two projects and user",
                    "remaining_datasets": Dataset.objects.count(),
                }
            ),
            flush=True,
        )


with TemporaryDirectory(prefix="workshop-mcp-acceptance-") as directory:
    replay = Replay(directory)
    try:
        replay.check("catalogue_auth_and_all_tool_validation", replay.discovery)
        replay.check("native_authored_semantic_version", replay.draft_and_semantic_import)
        replay.check("boundaries_isolation_and_failure_preservation", replay.boundaries)
        for size in SIZES:
            replay.check(
                f"large_file_pipeline_and_external_import_{size}",
                lambda size=size: replay.large(size),
            )
        replay.check("exploration_derived_chain_and_grouped_partition", replay.explore_partition)
        replay.check("trace_and_llm_call_sources", replay.trace_sources)
        replay.check("typed_targets_and_stepwise_cells", replay.typed_and_stepwise)
        replay.check("operation_values_and_atomic_failures", replay.operation_matrix)
        replay.check("concurrent_first_submissions", replay.concurrent_first_submissions)
        replay.check("all_parent_lineage_and_atomic_imports", replay.all_parent_lineage)
        replay.check("published_navigation_and_exact_query_pages", replay.navigation_and_query)
        replay.check("measured_exploration_and_independent_derivation", replay.exploration_results)
        replay.check("partition_contents_holdouts_and_retry", replay.partition_contents_and_retry)
        replay.check("complete_history_pagination", replay.full_history)
        replay.check(
            "every_workshop_tool_authorization", replay.permissions_for_every_workshop_tool
        )
        replay.check("cancel_running_worker_and_recover", replay.cancel_actual_worker)
        replay.check("tabular_file_upload_transform_export", replay.file_transport_formats)
        replay.check("source_bindings_artifacts_and_metadata", replay.source_bindings_and_metadata)
        if os.environ.get("WORKSHOP_ACCEPTANCE_DOCUMENTS"):
            replay.check("document_and_ocr_evidence_round_trip", replay.document_source_evidence)
        if os.environ.get("WORKSHOP_ACCEPTANCE_PDFS"):
            pdf_directory = Path(os.environ["WORKSHOP_ACCEPTANCE_PDFS"])
            for fixture in json.loads((pdf_directory / "manifest.json").read_text()):
                if fixture["kind"] != "batch":
                    replay.check(
                        "pdf_" + fixture["filename"].removesuffix(".pdf"),
                        lambda fixture=fixture: replay.pdf_case(fixture),
                    )
            for count in (10, 25, 100):
                replay.check(f"pdf_batch_{count}", lambda count=count: replay.pdf_quantity(count))
            replay.check("pdf_mixed_batch", replay.pdf_mixed_batch)
            replay.check("pdf_atomic_batch_recovery", replay.pdf_atomic_batch_recovery)
            replay.check("pdf_batch_limit_101", replay.pdf_batch_limit)
        replay.check("partition_maximum_name_lengths", replay.partition_name_boundary)
        replay.check("partition_declared_and_projected_groups", replay.partition_declared_groups)
        replay.check("partition_projected_group_aliases", replay.partition_group_aliases)
        replay.check(
            "partition_content_column_groups", lambda: replay.partition_declared_groups("content")
        )
        replay.check(
            "account_key_mutations_and_resource_links", replay.account_mutations_and_resources
        )
        replay.check("concurrent_partition_keys_across_sources", replay.concurrent_partition_keys)
        replay.check("catalogue_sampling_and_source_inventory", replay.catalogue_capability_audit)
        replay.check("large_partition_exact_membership", replay.large_partition_exact_membership)
        charged = BillingTelemetry.objects.filter(project=replay.project, amount__lt=0).count()
        assert charged == 0, "Workshop replay created unexpected billable usage"
        missing = sorted(EXPECTED - replay.called)
        print(
            json.dumps(
                {
                    "summary": {
                        "passed": sum(row["passed"] for row in replay.results),
                        "total": len(replay.results),
                        "successful_tools": sorted(replay.called),
                        "successful_tool_calls": dict(sorted(replay.success_counts.items())),
                        "rejected_tool_calls": dict(sorted(replay.error_counts.items())),
                        "resources_read": sorted(replay.resources),
                        "missing_workshop_tools": missing,
                        "max_response_bytes": replay.max_response,
                        "billable_usage_records": charged,
                        "rpc_performance": replay.rpc_performance(),
                    }
                }
            ),
            flush=True,
        )
    finally:
        replay.cleanup()
    assert all(row["passed"] for row in replay.results) and (
        not missing or os.environ.get("WORKSHOP_ACCEPTANCE_CASES")
    ), "Workshop acceptance failures remain"
