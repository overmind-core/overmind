"""Writes the document fixtures: `uv run python tests/fixtures/documents/generate.py`."""

from __future__ import annotations

import hashlib
import struct
from pathlib import Path

import docx
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).parent
PAD = bytes.fromhex("28BF4E5E4E758A4164004E56FFFA01082E2E00B6D0683E802F0CA9FE6453697A")


def _rc4(key: bytes, data: bytes) -> bytes:
    state = list(range(256))
    j = 0
    for i in range(256):
        j = (j + state[i] + key[i % len(key)]) % 256
        state[i], state[j] = state[j], state[i]
    out = bytearray()
    i = j = 0
    for byte in data:
        i = (i + 1) % 256
        j = (j + state[i]) % 256
        state[i], state[j] = state[j], state[i]
        out.append(byte ^ state[(state[i] + state[j]) % 256])
    return bytes(out)


def pdf(pages: list[str], *, password: str | None = None) -> bytes:
    """A minimal PDF with one Helvetica text line per page, optionally RC4-40 encrypted
    (standard security handler, revision 2) with a user password."""
    file_id = hashlib.md5(b"overmind-fixture").digest()
    key = b""
    encrypt = ""
    if password is not None:
        owner = _rc4(
            hashlib.md5((b"owner" + PAD)[:32]).digest()[:5], (password.encode() + PAD)[:32]
        )
        permissions = struct.pack("<i", -4)
        key = hashlib.md5((password.encode() + PAD)[:32] + owner + permissions + file_id).digest()[
            :5
        ]
        user = _rc4(key, PAD)
        encrypt = f"<< /Filter /Standard /V 1 /R 2 /O <{owner.hex()}> /U <{user.hex()}> /P -4 >>"
    count = len(pages)
    objects = {
        1: "<< /Type /Catalog /Pages 2 0 R >>",
        2: "<< /Type /Pages /Kids ["
        + " ".join(f"{4 + 2 * i} 0 R" for i in range(count))
        + f"] /Count {count} >>",
        3: "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    streams = {}
    for i, text in enumerate(pages):
        page, content = 4 + 2 * i, 5 + 2 * i
        objects[page] = (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content} 0 R >>"
        )
        body = f"BT /F1 18 Tf 72 700 Td ({text}) Tj ET".encode() if text else b""
        if key:
            object_key = hashlib.md5(key + struct.pack("<i", content)[:3] + b"\0\0").digest()[:10]
            body = _rc4(object_key, body)
        streams[content] = body
    if encrypt:
        objects[len(objects) + len(streams) + 1] = encrypt
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for number in sorted({*objects, *streams}):
        offsets[number] = len(out)
        if number in streams:
            body = streams[number]
            out += f"{number} 0 obj\n<< /Length {len(body)} >>\nstream\n".encode()
            out += body + b"\nendstream\nendobj\n"
        else:
            out += f"{number} 0 obj\n{objects[number]}\nendobj\n".encode()
    xref = len(out)
    size = max(offsets) + 1
    out += f"xref\n0 {size}\n0000000000 65535 f \n".encode()
    for number in range(1, size):
        out += f"{offsets[number]:010d} 00000 n \n".encode()
    trailer = f"/Size {size} /Root 1 0 R /ID [<{file_id.hex()}> <{file_id.hex()}>]"
    if encrypt:
        trailer += f" /Encrypt {max(offsets)} 0 R"
    out += f"trailer\n<< {trailer} >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def _picture(lines: list[str]) -> Image.Image:
    image = Image.new("RGB", (1600, 220 + 120 * len(lines)), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=72)
    for i, line in enumerate(lines):
        draw.text((100, 100 + 120 * i), line, fill="black", font=font)
    return image


def main() -> None:
    (HERE / "native.pdf").write_bytes(
        pdf(["Invoice total: 1234.56 USD", "Refunds are issued within 30 days."])
    )
    (HERE / "password.pdf").write_bytes(pdf(["Secret terms"], password="secret"))
    (HERE / "blank.pdf").write_bytes(pdf([""]))
    _picture(["Policy 7: refunds need a receipt."]).save(
        HERE / "scanned.pdf", "PDF", resolution=200
    )
    _picture(["Order 991 shipped on Monday"]).save(HERE / "text.png")
    Image.new("RGB", (800, 400), "white").save(HERE / "blank.png")
    document = docx.Document()
    document.add_heading("Returns policy", level=1)
    document.add_paragraph("Customers may return items within 30 days.")
    table = document.add_table(rows=2, cols=2)
    for row, cells in zip(table.rows, (("Item", "Days"), ("Shoes", "30")), strict=True):
        for cell, value in zip(row.cells, cells, strict=True):
            cell.text = value
    document.save(HERE / "policy.docx")
    (HERE / "notes.md").write_text(
        "# Shipping\n\nOrders ship in two days.\n\n| Zone | Days |\n| ---- | ---- |\n| EU   | 3    |\n",
        encoding="utf-8",
    )
    (HERE / "faq.txt").write_text(
        "How do I reset my password?\nUse the reset link.\n\n\n\nWhere is my order?\n",
        encoding="utf-8",
    )
    (HERE / "latin1.txt").write_bytes("Café crème\n".encode("latin-1"))


if __name__ == "__main__":
    main()
